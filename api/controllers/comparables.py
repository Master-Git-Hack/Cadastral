from base64 import b64encode
from datetime import datetime, timedelta
from time import sleep
from typing import Any, Dict

from babel.dates import format_date
from dateparser import parse
from dateutil.relativedelta import relativedelta
from num2words import num2words
from openpyxl import Workbook

# from openpyxl.drawing.image import Image
from openpyxl.styles import Border, Font, PatternFill, Side
from PIL import Image
from requests import get
from sqlalchemy.orm import Session
from sqlalchemy.sql import func

from .. import config, database
from ..controllers.comparables_catcom import ComparablesCatComReport
from ..models.cedula_comparables import CedulaComparables
from ..models.cedula_mercado import CedulaMercado
from ..models.comparables_catcom import ComparablesCatCom
from ..utils.locale import as_complete_date, as_currency


class CedulaService:
    text = dict(
        header=dict(
            black=Font(name="Arial", size=18, bold=True),
            white=Font(name="Arial", size=18, bold=True, color="FFFFFF"),
        ),
        title=dict(
            black=Font(name="Arial", size=14, bold=True),
            white=Font(name="Arial", size=14, bold=True, color="FFFFFF"),
        ),
        sub_title=dict(
            black=Font(name="Arial", size=12, bold=True),
            white=Font(name="Arial", size=12, bold=True, color="FFFFFF"),
        ),
        normal=dict(
            black=Font(name="Arial", size=9),
            white=Font(name="Arial", size=9, color="FFFFFF"),
        ),
        bold=dict(
            black=Font(name="Arial", size=9, bold=True),
            white=Font(name="Arial", size=9, bold=True, color="FFFFFF"),
        ),
    )
    background = dict(
        red=PatternFill(start_color="fd0d00", end_color="fd0d00", fill_type="solid"),
        green=PatternFill(start_color="10a870", end_color="10a870", fill_type="solid"),
        violet=PatternFill(start_color="b2a1c7", end_color="b2a1c7", fill_type="solid"),
        kaki=PatternFill(start_color="ddd9c3", end_color="ddd9c3", fill_type="solid"),
        teal=PatternFill(start_color="12b050", end_color="12b050", fill_type="solid"),
        blue=PatternFill(start_color="548dd4", end_color="548dd4", fill_type="solid"),
        pink=PatternFill(start_color="d99594", end_color="d99594", fill_type="solid"),
        emerald=PatternFill(
            start_color="10a870", end_color="10a870", fill_type="solid"
        ),
        yellow=PatternFill(start_color="fffd03", end_color="fffd03", fill_type="solid"),
    )
    border = dict(
        top=Border(top=Side(border_style="thin", color="000000")),
        bottom=Border(bottom=Side(border_style="thin", color="000000")),
        left=Border(left=Side(border_style="thin", color="000000")),
        right=Border(right=Side(border_style="thin", color="000000")),
        x=Border(
            left=Side(border_style="thin", color="000000"),
            right=Side(border_style="thin", color="000000"),
        ),
        y=Border(
            top=Side(border_style="thin", color="000000"),
            bottom=Side(border_style="thin", color="000000"),
        ),
        full=Border(
            left=Side(border_style="thin", color="000000"),
            right=Side(border_style="thin", color="000000"),
            top=Side(border_style="thin", color="000000"),
            bottom=Side(border_style="thin", color="000000"),
        ),
    )
    mercado_headers_template = [
        tipo,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        "$ DLLS",
        None,
        None,
        None,
        None,
        None,
        None,
    ]
    mercado_headers_template2 = [
        "Datos de Verificación",
        None,
        None,
        None,
        None,
        None,
        None,
        "Ubicación",
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        "Características de Terreno",
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        "Características de Construcción",
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        "Infraestructura",
        None,
        "Valores",
        None,
        None,
        None,
        None,
        None,
        "Vigencia",
        None,
        None,
        None,
        "Elaboró",
    ]
    mercado_headers_template3 = [
        "Folio",
        "Tipo de Inmueble",
        "Tipo de Operación",
        "Fecha de Captura",
        "URL Fuente",
        "Informante",
        "Teléfono de Informante",
        "Coordenada UTM X",
        "Coordenada UTM Y",
        "Estado",
        "Municipio",
        "Ciudad o Población",
        "Tipo y Nombre de Colonia / Asentamiento",
        "Tipo y Nombre de la Calle",
        "No. Exterior",
        "No. Interior",
        "Nombre Edificio / Proto / Predio",
        "Régimen de Propiedad",
        "Clasificación Periférica",
        "Clasificación Económica de la Zona (Campo)",
        "Uso de Suelo Carta, Uso Plan",
        "Entre Calles",
        "Ubicación en la Manzana",
        "Número de Frentes",
        "Superficie Terreno M2",
        "Frente ML",
        "Frente Tipo ML",
        "Fondo",
        "Forma",
        "Topografía",
        "Superficie Construcción M2",
        "Proyecto",
        "Edo. Conservación",
        "Tipo de Construcción",
        "Calidad",
        "Edad",
        "Niveles",
        "Unidades Rentables",
        "Descripción de Espacios",
        "T / C",
        "Servicios",
        "Descripción de Servicios",
        "Precio",
        "Precio Unitario",
        "Precio Total USD",
        "Precio Unitario USD",
        "Precio Total Aplicable en la Homologación MXN",
        "Precio Unitario Aplicable en la Homologación MXN",
        "Observaciones",
        "Hoy",
        "Días",
        "Caduca en 6 meses Fecha",
        "Elaboró",
    ]
    mercado_width = {
        "A": 92,
        "B": 183,
        "C": 90,
        "D": 186,
        "E": 90,
        "F": 237,
        "G": 94,
        "H": 92,
        "I": 92,
        "J": 108,
        "K": 209,
        "L": 90,
        "M": 208,
        "N": 136,
        "O": 90,
        "P": 90,
        "Q": 103,
        "R": 136,
        "S": 94,
        "T": 239,
        "U": 215,
        "V": 204,
        "W": 150,
        "X": 76,
        "Y": 101,
        "Z": 76,
        "AA": 76,
        "AB": 76,
        "AC": 125,
        "AD": 162,
        "AE": 105,
        "AF": 90,
        "AG": 149,
        "AH": 149,
        "AI": 104,
        "AJ": 90,
        "AK": 90,
        "AL": 90,
        "AM": 183,
        "AN": 90,
        "AO": 143,
        "AP": 543,
        "AQ": 120,
        "AR": 120,
        "AS": 120,
        "AT": 120,
        "AU": 120,
        "AV": 120,
        "AW": 246,
        "AX": 79,
        "AY": 72,
        "AZ": 183,
        "BA": 77,
    }

    def __init__(self, db: Session, user=None):
        self.db = db
        self.user = user
        self.comparable = Comparables(db)
        self.cedula = Cedula(db)
        self.mercado = Mercado(db)

    def get_cedulas(self):
        cedula = self.mercado
        results = cedula.filter_group(usuario=self.user.usuario)
        if results is None:
            return {"message": "No se encontraron cedulas", "status_code": 404}
        return {
            "data": [
                cedula.dict(data)
                for data in cedula.current
                if data.fecha + timedelta(days=180) > datetime.now()
            ]
        }

    def get_cedula_by_id(self, id: int):
        cedula = self.mercado
        result = cedula.filter(id=id, usuario=self.user.usuario)
        return (
            {"data": cedula.dict()}
            if result
            else {"message": "No se encontraron cedulas", "status_code": 404}
        )

    def create_cedula(self, registro: str):
        cedula = self.mercado
        result = cedula.create(registro=registro, usuario=self.user.usuario)
        return (
            {"data": cedula.dict()}
            if result
            else {"message": "No se pudo crear la cedula", "status_code": 404}
        )

    def update_cedula(self, id: int, data: Dict[str, Any]):
        cedula = self.mercado
        if not cedula.filter(id=id, usuario=self.user.usuario):
            return {"message": "No se pudo encontrar la cedula", "status_code": 404}
        return (
            {"data": cedula.dict()}
            if cedula.update(**data)
            else {"message": "No se pudo actualizar la cedula", "status_code": 404}
        )

    def delete_cedula(self, id: int):
        cedula = self.mercado
        if not cedula.filter(id=id):
            return {"message": "No se pudo encontrar la cedula", "status_code": 404}
        return (
            {"data": cedula.dict()}
            if cedula.delete(id)
            else {"message": "No se pudo eliminar la cedula", "status_code": 404}
        )

    def get_comparables(self, cedula_mercado: int):
        comp = self.cedula
        results = comp.filter_group(id_cedula_mercado=cedula_mercado)
        return (
            {"data": comp.list()}
            if results
            else {"message": "No se encontraron comparables", "status_code": 404}
        )

    def get_comparable_by_id(self, id: int):
        comp = self.cedula
        result = comp.get(id)
        return (
            {"data": comp.dict()}
            if result
            else {"message": "No se encontraron comparables", "status_code": 404}
        )

    def create_comparable(
        self, cedula_mercado: int, tipo: str, comparable: int, data: Dict[str, Any]
    ):
        comp = self.cedula
        if comp.filter(
            id_cedula_mercado=cedula_mercado, id_comparable_catcom=comparable, tipo=tipo
        ):
            return {"message": "Ya existe el comparable", "status_code": 404}
        return (
            {"data": comp.dict()}
            if comp.create(
                tipo=tipo.upper(),
                id_cedula_mercado=cedula_mercado,
                id_comparable_catcom=comparable,
                **data,
            )
            else {"message": "No se pudo crear el comparable", "status_code": 404}
        )

    def update_comparable(self, id: int, data: Dict[str, Any]):
        comp = self.cedula
        if not comp.get(id):
            return {"message": "No se encontraron comparables", "status_code": 404}
        return (
            {"data": comp.dict()}
            if comp.update(**data)
            else {"message": "No se pudo actualizar el comparable", "status_code": 404}
        )

    def delete_comparable(self, id: int):
        comp = self.cedula
        if not comp.get(id):
            return {"message": "No se encontraron comparables", "status_code": 404}
        try:
            return (
                {"data": comp.dict()}
                if comp.delete(id)
                else {
                    "message": "No se pudo eliminar el comparable",
                    "status_code": 404,
                }
            )
        except Exception:
            sleep(5)
            return (
                {"data": comp.dict()}
                if comp.delete(id)
                else {
                    "message": "No se pudieron eliminar los comparables",
                    "status_code": 404,
                }
            )

    @staticmethod
    def check_services(**kwargs):
        servicios = [
            "agua",
            "drenaje",
            "energia_electrica",
            "alumbrado_publico",
            "banqueta",
            "pavimento",
            "telefonia",
        ]
        completos = all(kwargs.get(servicio, False) for servicio in servicios)
        return (
            "Sí Tiene Completos"
            if completos
            else "Tiene Algunos"
            if any(kwargs.values())
            else "No Tiene"
        )

    @staticmethod
    def check_desc_services(**kwargs):
        servicios = {
            "agua": "Red de Agua Potable",
            "drenaje": "Red de Drenaje",
            "energia_electrica": "Red de Energía Eléctrica",
            "alumbrado_publico": "Alumbrado Público, Voz y Datos",
            "pavimento": "Pavimento",
            "banqueta": "Banquetas",
        }
        return ", ".join(
            valor for clave, valor in servicios.items() if kwargs.get(clave, False)
        )

    @staticmethod
    def get_image(image_url: str):
        try:
            response = get(image_url)
            response.raise_for_status()
            return b64encode(response.content).decode("utf-8")
        except Exception:
            return ""

    def generate_preview(self, cedula_mercado: int, data: Dict[str, Any]):
        if not self.mercado.get(cedula_mercado) or not self.cedula.filter_group(
            id_cedula_mercado=cedula_mercado
        ):
            return {"message": "No se encontraron comparables", "status_code": 404}

        mercado_data = self.mercado.dict(excludes=["id", "usuario", "fecha"])
        cedulas = [
            {**self.comp.dict(), "tipo": c.tipo, **mercado_data}
            for c in self.cedula.current
            if self.comp.get(c.id_comparable_catcom)
        ]

        if "ids" in data:
            cedulas = [c for c in cedulas if c.get("id") in data["ids"]]

        grouped_data = {
            tipo: [c for c in cedulas if c.get("tipo") == tipo]
            for tipo in {c.get("tipo") for c in cedulas}
        }

        for records in grouped_data.values():
            for record in records:
                record["fecha_captura"] = as_complete_date(record["fecha_captura"])

        return {
            "data": [
                {"tipo": tipo, "records": records}
                for tipo, records in grouped_data.items()
            ]
        }

    def _fetch_image(self, filename: str):
        url_base = "http://172.31.113.151/comparables/imagenes"
        image_url = f"{url_base}/{filename}"
        try:
            response = get(image_url)
            path = f"{config.PATHS.TMP}/{filename}"
            with open(path, "wb") as f:
                f.write(response.content)
            return path
        except Exception:
            return None

    def handle_images(self, filename: str, crop: bool = False):
        if not filename or filename == "null":
            return {"message": "No se encontró la imagen", "status_code": 404}

        path = self._fetch_image(filename)
        if not path:
            return {"message": "No se pudo descargar la imagen", "status_code": 404}

        ext = filename.split(".")[-1]
        if crop:
            img = Image.open(path)
            width, height = img.size
            left, top = (width - 200) // 2, (height - 200) // 2
            right, bottom = (width + 200) // 2, (height + 200) // 2
            cropped = img.crop((left, top, right, bottom))
            path = path.replace(f".{ext}", f"_crop.{ext}")
            cropped.resize((width, height), Image.Resampling.LANCZOS)
            cropped.save(path, quality=95)

        return {
            "filename": filename,
            "path": path,
            "media_type": f"image/{ext}",
            "delete": True,
        }

    def handle_image_preview(
        self, comparable: int, width: int = 200, height: int = 200
    ):
        if not (comp_data := self.comp.get(comparable)):
            return {"message": "Registro No encontrado"}

        filename = comp_data.imagen_2
        if not filename:
            return {"message": "Imagen no encontrada"}

        path = self._fetch_image(filename)
        if not path:
            return {"message": "No se pudo descargar la imagen", "status_code": 404}

        ext = filename.split(".")[-1]
        img = Image.open(path)
        img_width, img_height = img.size
        left, top = (img_width - width) // 2, (img_height - height) // 2
        right, bottom = (img_width + width) // 2, (img_height + height) // 2
        cropped = img.crop((left, top, right, bottom))
        path = path.replace(f".{ext}", f"_crop.{ext}")
        cropped.resize((img_width, img_height), Image.Resampling.LANCZOS)
        cropped.save(path, quality=95)

        return {
            "filename": filename,
            "path": path,
            "media_type": f"image/{ext}",
            "delete": True,
        }

    def generate_xlsx(self, cedula_mercado: int, data: Dict[str, Any], db: Session):
        comp = Comparables(db)
        cedula = Cedula(db)
        mercado = Mercado(db)

        if mercado.get(cedula_mercado) is None:
            return dict(message="No se encontraron comparables", status_code=404)
        if len(cedula.filter_group(id_cedula_mercado=cedula_mercado)) == 0:
            return dict(message="No se encontraron comparables", status_code=404)

        registro = mercado.current.registro
        mercado_data = mercado.dict(excludes=["id", "usuario", "fecha"])

        # Pre-filtered cedulas
        cedulas = [
            {**comp.dict(), "tipo": c.tipo, **mercado_data}
            for c in cedula.current
            if comp.get(c.id_comparable_catcom) is not None
        ]

        # Handle missing ids gracefully
        ids = data.get("ids", [])
        if not ids:
            return dict(message="No se encontraron comparables", status_code=404)

        # Filter cedulas by provided ids
        filtered_cedulas = [c for c in cedulas if c.get("id") in ids]

        # Split cedulas by tipo
        tipos = {"TERRENO": [], "VENTA": [], "RENTA": []}
        for c in filtered_cedulas:
            tipo = c.get("tipo", "").upper()
            if tipo in tipos:
                tipos[tipo].append(c)

        # Create the workbook and configure the sheet
        workbook = Workbook()
        mercado_sheet = workbook.active
        mercado_sheet.title = "MERCADO"
        url_base = "http://172.31.113.151/comparables/imagenes"

        # Define common cell styles
        header_font = self.text["header"]["white"]
        sub_title_font = self.text["sub_title"]["black"]
        bold_font = self.text["bold"]["black"]
        normal_font = self.text["normal"]["black"]
        fill_colors = {
            "TERRENO": self.background["red"],
            "VENTA": self.background["green"],
            "RENTA": self.background["violet"],
        }

        # Function to format URLs
        def format_url(image_key):
            return f"{url_base}/{image_key}" if image_key else ""

        # Function to calculate price per unit
        def calculate_price_per_unit(r, field="valor_total_mercado"):
            try:
                return r.get(field, 0) / r.get("superficie_terreno", 1)
            except:
                return 0

        # Iterate over the tipos (TERRENO, VENTA, RENTA)
        for tipo, cedulas_list in tipos.items():
            if not cedulas_list:
                continue

            # Add header for the type
            mercado_sheet.append(self.mercado_headers_template)
            mercado_sheet.merge_cells(
                start_row=mercado_sheet.max_row,
                start_column=1,
                end_row=mercado_sheet.max_row,
                end_column=4,
            )
            cell = mercado_sheet[
                mercado_sheet.cell(row=mercado_sheet.max_row, column=1).coordinate
            ]
            cell.font = header_font
            cell.fill = fill_colors[tipo]

            # Subheader
            mercado_sheet.append(self.mercado_headers_template2)
            mercado_last_row = mercado_sheet.max_row
            mercado_sheet.merge_cells(f"A{mercado_last_row}:G{mercado_last_row}")
            cell = mercado_sheet[
                mercado_sheet.cell(row=mercado_last_row, column=1).coordinate
            ]
            cell.font = sub_title_font
            cell.fill = self.background["kaki"]

            # Additional sections and formatting
            self._format_additional_sections(mercado_sheet, mercado_last_row)

            # Add data rows
            for index, r in enumerate(cedulas_list):
                r["imagen_1"] = format_url(r.get("imagen_1"))
                r["imagen_2"] = format_url(r.get("imagen_2"))
                r["imagen_3"] = format_url(r.get("imagen_3"))
                r["captura_pantalla"] = format_url(r.get("captura_pantalla"))

                r["fecha_captura"] = as_complete_date(r.get("fecha_captura", "hoy"))
                date = parse(r.get("fecha_captura"))
                hoy = datetime.now()
                dias = (hoy - date).days
                hoy = format_date(hoy, format="d 'de' MMMM 'del' y", locale="es")
                seis_meses = format_date(
                    date + relativedelta(months=6),
                    format="d 'de' MMMM 'del' y",
                    locale="es",
                )
                numero_frentes = r.get("numero_frentes", 1)
                numero_frentes = (
                    f"{numero_frentes} ({num2words(numero_frentes, lang='es').upper()})"
                )

                # Calculate price per unit
                precio_unitario = calculate_price_per_unit(r)
                precio_unitario_usd = calculate_price_per_unit(r, "vtm_usd")

                # Append data to the sheet
                mercado_sheet.append(
                    [
                        index + 1,
                        r.get("tipo_inmueble", "N/A"),
                        r.get("tipo_operacion", "N/A"),
                        r.get("fecha_captura", "N/A"),
                        r.get("url_fuente", "N/A"),
                        r.get("nombre_anunciante", "N/A"),
                        r.get("telefono_anunciante", "N/A"),
                        r.get("x_utm", 0),
                        r.get("y_utm", 0),
                        r.get("estado", "GUANAJUATO"),
                        r.get("municipio", "GUANAJUATO"),
                        r.get("localidad", "GUANAJUATO"),
                        f"{r.get('tipo_asentamiento')} {r.get('nombre_asentamiento')}",
                        f"{r.get('tipo_vialidad')} {r.get('nombre_vialidad')}",
                        r.get("numero_exterior"),
                        r.get("numero_interior"),
                        r.get("edificio"),
                        r.get("regimen_propiedad"),
                        r.get("tipo_zona"),
                        r.get("uso_suelo_observado"),
                        r.get("uso_suelo_oficial"),
                        r.get("entrecalles"),
                        r.get("ubicacion_manzana"),
                        numero_frentes,
                        r.get("superficie_terreno"),
                        r.get("longitud_frente"),
                        r.get("longitud_frente_tipo"),
                        r.get("longitud_fondo"),
                        r.get("forma"),
                        r.get("topografia"),
                        r.get("superficie_construccion"),
                        r.get("calidad_proyecto"),
                        r.get("estado_conservacion"),
                        r.get("tipo_construccion"),
                        r.get("calidad_construccion"),
                        r.get("edad"),
                        r.get("niveles"),
                        r.get("unidades_rentables"),
                        r.get("descripcion_espacios"),
                        tc,
                        self.check_services(**r),
                        self.check_desc_services(**r),
                        as_currency(r.get("valor_total_mercado", 0), "$ -"),
                        as_currency(precio_unitario, "$ -"),
                        as_currency(r.get("vtm_usd", 0), "$ -"),
                        as_currency(precio_unitario_usd, "$ -"),
                        as_currency(0, "$ -"),
                        as_currency(0, "$ -"),
                        r.get("observaciones"),
                        hoy,
                        dias,
                        seis_meses,
                        r.get("usuario"),
                    ]
                )

                # Set font for cells in the row
                for i in range(1, 53):
                    current_cell = mercado_sheet.cell(
                        row=mercado_sheet.max_row, column=i
                    )
                    current_cell.font = normal_font

            # Append an empty row for separation
            mercado_sheet.append([])

        # Set column widths
        for column, width in self.mercado_width.items():
            mercado_sheet.column_dimensions[column].width = width / 7.5

        filename = f"{registro}.xlsx"
        path = f"{config.PATHS.TMP}/{filename}"
        workbook.save(path)

        return dict(filename=filename, path=path, delete=True)

    def _format_additional_sections(self, sheet, last_row):
        """Helper function to format additional sections."""
        sections = [
            ("H", "Q", self.background["violet"]),
            ("S", "AD", self.background["red"]),
            ("AE", "AN", self.background["teal"]),
            ("AO", "AP", self.background["kaki"]),
            ("AQ", "AV", self.background["blue"]),
            ("AW", "AZ", self.background["pink"]),
        ]

        for start_col, end_col, color in sections:
            sheet.merge_cells(f"{start_col}{last_row}:{end_col}{last_row}")
            for col in range(ord(start_col), ord(end_col) + 1):
                cell = sheet.cell(row=last_row, column=col - 64)
                cell.fill = color
